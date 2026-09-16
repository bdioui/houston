"""Lecture d'un export SIFAC (XLSX) vers des lignes prêtes pour `SifacLine`.

Port de `frontend/src/lib/sifac/parse.ts`. Fonction pure : elle ne touche ni la
base ni le tenant, elle lit un fichier et rend des dictionnaires.

Deux écarts délibérés avec l'original TypeScript, tous deux dus au fait que la
destination n'est plus Grist mais une colonne Django typée :

- une date absente vaut `None` et non `''` — un `DateField` refuse la chaîne
  vide ;
- les montants sont des `Decimal` et non des `float`, parce que la colonne est
  un `DecimalField` et que l'agrégation compare des totaux au centime près.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

# SIFAC tronque ses en-têtes à 30 caractères ("Montant réceptionné non factur")
# et emploie des apostrophes typographiques : l'appariement se fait sur un
# préfixe normalisé, jamais sur l'égalité.
HEADERS: dict[str, str] = {
    "pfi": "programme de financement",
    "flux_id": "numero de flux",
    "flux_label": "libelle du flux",
    "rubrique": "rubrique de la piece",
    "supplier_name": "nom du tiers",
    "supplier_code": "numero du tiers fournisseur",
    "account": "compte general",
    "account_label": "libelle compte general",
    "engagement_date": "date initiale",
    "amount_engaged": "montant engage htr",
    "amount_certified": "montant htr des sf",
    "amount_received": "montant receptionne",
    "csf_date": "date comptable du csf",
    "invoice_number": "numero de facture",
    "invoice_date": "date comptable facture",
    "invoice_text": "texte facture",
    "amount_invoiced": "montant facture htr",
    "amount_paid": "montant paye",
    "payment_date": "date de paiement",
    "amount_report": "report",
    "otp": "element d'otp",
    "category": "compte d'execution budgetaire",
}

TEXT_FIELDS = (
    "pfi", "flux_id", "flux_label", "rubrique", "supplier_name",
    "supplier_code", "account", "account_label", "invoice_number",
    "invoice_text", "otp", "category",
)
DATE_FIELDS = ("engagement_date", "csf_date", "invoice_date", "payment_date")
AMOUNT_FIELDS = (
    "amount_engaged", "amount_certified", "amount_received",
    "amount_invoiced", "amount_paid", "amount_report",
)


class SifacParseError(ValueError):
    """Erreur de lecture imputable au fichier, pas au code.

    Distincte de ValueError pour que la vue puisse la traduire en 400 sans
    avaler au passage les vraies erreurs de programmation.
    """


def normalize_header(v: object) -> str:
    s = "" if v is None else str(v)
    # NFD sépare les diacritiques combinants, la classe Mn les retire.
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")
    s = s.replace("‘", "'").replace("’", "'")
    # \s couvre l'espace insécable, dont SIFAC est friand.
    s = re.sub(r"\s+", " ", s)
    return s.strip().lower()


def to_text(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def to_amount(v: object) -> Decimal:
    """Rend toujours un Decimal à deux décimales, 0 si la valeur est illisible.

    Le silence est volontaire et hérité du TypeScript : une cellule vide ou un
    tiret dans une colonne de montant est courant dans les exports SIFAC, et
    vaut zéro. Ce sont les en-têtes, pas les cellules, dont l'absence est
    traitée comme une erreur.
    """
    if isinstance(v, bool):
        return Decimal("0.00")
    if isinstance(v, (int, float, Decimal)):
        try:
            return Decimal(str(v)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return Decimal("0.00")
    if isinstance(v, str):
        cleaned = re.sub(r"\s", "", v).replace(",", ".")
        try:
            return Decimal(cleaned).quantize(Decimal("0.01"))
        except InvalidOperation:
            return Decimal("0.00")
    return Decimal("0.00")


def to_date(v: object) -> date | None:
    """`None` et non `''` : la destination est un DateField.

    Pas de conversion de fuseau. Le TypeScript devait lire les composantes
    locales d'un objet Date pour éviter qu'un `toISOString()` ne recule la date
    d'un jour ; ici openpyxl rend un datetime naïf, qui ne porte pas le piège.
    """
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", v.strip())
        if m:
            try:
                return date(int(m[1]), int(m[2]), int(m[3]))
            except ValueError:
                return None
    return None


def map_headers(header_row: list[object]) -> dict[str, int]:
    """Une colonne mal appariée serait silencieusement vide en base : on échoue
    ici, bruyamment, plutôt que de laisser passer un import tronqué.
    """
    normalized = [normalize_header(c) for c in header_row]
    index: dict[str, int] = {}
    missing: list[str] = []

    for field, prefix in HEADERS.items():
        matches = [i for i, h in enumerate(normalized) if h.startswith(prefix)]
        if not matches:
            missing.append(f'{field} ("{prefix}…")')
        elif len(matches) > 1:
            labels = ", ".join(f'"{header_row[i]}"' for i in matches)
            raise SifacParseError(
                f"Colonne SIFAC ambiguë pour {field} : {labels}."
            )
        else:
            index[field] = matches[0]

    if missing:
        raise SifacParseError(
            f"Colonnes SIFAC introuvables : {', '.join(missing)}."
        )
    return index


def suggest_exercice(rows: list[dict]) -> int:
    """L'exercice n'est pas une donnée du fichier : les lignes de report portent
    la date d'engagement de l'année précédente. On propose l'année majoritaire,
    à charge pour l'utilisateur de confirmer.
    """
    years: dict[int, int] = {}
    for r in rows:
        d = r["invoice_date"] or r["payment_date"] or r["engagement_date"]
        if d is None:
            continue
        years[d.year] = years.get(d.year, 0) + 1
    if not years:
        raise SifacParseError(
            "Aucune date exploitable : exercice indéterminable."
        )
    # Départage par année décroissante à égalité de compte, pour que deux
    # lectures du même fichier proposent toujours le même exercice — un dict
    # Python conserve l'ordre d'insertion, qui dépendrait sinon du tri des
    # lignes dans l'export.
    return max(years.items(), key=lambda kv: (kv[1], kv[0]))[0]


def parse_sifac_export(source) -> tuple[str, int, list[dict]]:
    """Rend (pfi, exercice proposé, lignes). `source` est un fichier ou un chemin.

    Les lignes rendues portent déjà `exercice`, mais c'est la valeur proposée :
    l'appelant la remplace par celle que l'utilisateur confirme.
    """
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        # openpyxl ne lit que le XLSX. Un vrai .xls (BIFF, antérieur à 2007)
        # échoue ici, comme il échouait déjà côté navigateur.
        raise SifacParseError(
            "Fichier illisible : un export SIFAC au format .xlsx est attendu."
        ) from exc

    try:
        sheet = wb[wb.sheetnames[0]]
        grid = [list(r) for r in sheet.iter_rows(values_only=True)]
    finally:
        wb.close()

    if len(grid) < 2:
        raise SifacParseError("Fichier SIFAC vide ou sans ligne d'en-tête.")

    header_row, data_rows = grid[0], grid[1:]
    index = map_headers(header_row)

    def at(row: list, field: str):
        i = index[field]
        return row[i] if i < len(row) else None

    parsed: list[dict] = []
    for row in data_rows:
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue
        r: dict = {f: to_text(at(row, f)) for f in TEXT_FIELDS}
        r.update({f: to_date(at(row, f)) for f in DATE_FIELDS})
        r.update({f: to_amount(at(row, f)) for f in AMOUNT_FIELDS})
        # Les lignes sans numéro de flux sont des sous-totaux : écartées ici,
        # jamais insérées.
        if r["flux_id"] == "":
            continue
        parsed.append(r)

    if not parsed:
        raise SifacParseError("Aucune ligne exploitable dans le fichier.")

    # Le périmètre d'écrasement est le couple (PFI, exercice) : un fichier
    # multi-PFI le rendrait indéfini et effacerait le mauvais programme.
    pfis = list(dict.fromkeys(r["pfi"] for r in parsed if r["pfi"]))
    if len(pfis) != 1:
        raise SifacParseError(
            f"L'export doit porter sur un seul PFI, {len(pfis)} trouvés : "
            f"{', '.join(pfis)}."
        )

    exercice = suggest_exercice(parsed)
    for r in parsed:
        r["exercice"] = exercice
    return pfis[0], exercice, parsed
